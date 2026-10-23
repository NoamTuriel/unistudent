"""Turning course documents into text pages. Best available tool first, pure-Python fallback.

Returns a Conversion: a list of page texts, the method used, and warnings.
Pages whose text layer is (almost) empty or OCR noise, and slides with a picture, are reported, so Claude can read
them visually.
"""
import logging
import re
import shutil
import subprocess
import tempfile
import unicodedata
import warnings
from dataclasses import dataclass, field
from pathlib import Path

logging.getLogger("pypdf").setLevel(logging.CRITICAL)

DOCUMENTS = {".pdf", ".docx", ".doc", ".odt", ".rtf", ".pptx", ".ppt", ".odp", ".txt", ".md",
             ".xlsx", ".xls", ".html", ".htm", ".epub"}
MARKITDOWN_ONLY = {".xlsx", ".xls", ".html", ".htm", ".epub"}
RECORDINGS = {".mp4", ".mkv", ".mov", ".m4v", ".webm", ".avi", ".mp3", ".m4a", ".wav", ".aac"}
IMAGES = {".png", ".jpg", ".jpeg", ".heic", ".webp", ".gif", ".tif", ".tiff"}
OFFICE = {".docx", ".doc", ".odt", ".rtf", ".pptx", ".ppt", ".odp"}
EMPTY_PAGE_CHARS = 20
NOT_PAGED = {"text", "markitdown", "python-docx"}  # one block of text, no pages to read visually


@dataclass
class Conversion:
    pages: list
    method: str
    warnings: list = field(default_factory=list)
    picture_pages: list = field(default_factory=list)  # slides whose picture (often a graph) has no text
    view: Path = None  # LibreOffice's PDF of an Office file, kept so a worker can look at a page (Read can't open .pptx)

    @property
    def empty_pages(self):
        """The pages to read visually: no text layer, or a slide with a picture."""
        if self.method in NOT_PAGED:  # a short text file is not a page with a missing text layer
            return []
        return sorted({i + 1 for i, text in enumerate(self.pages) if len(text.strip()) < EMPTY_PAGE_CHARS}
                      | set(self.picture_pages))


def kind(path) -> str:
    suffix = Path(path).suffix.lower()
    if suffix in DOCUMENTS:
        return "document"
    if suffix in RECORDINGS:
        return "recording"
    if suffix in IMAGES:
        return "image"
    return "other"


WORKS = {}  # program path → whether it runs: a wrapper left after LibreOffice was removed does not


def _soffice():
    path = shutil.which("soffice") or shutil.which("libreoffice")
    if path and path not in WORKS:
        try:
            WORKS[path] = subprocess.run([path, "--version"], capture_output=True, timeout=60).returncode == 0
        except (OSError, subprocess.TimeoutExpired):
            WORKS[path] = False
    return path if path and WORKS[path] else None


def _pdf_pages(path: Path) -> Conversion:
    if shutil.which("pdftotext"):
        result = subprocess.run(["pdftotext", "-layout", "-enc", "UTF-8", str(path), "-"],
                                capture_output=True)
        if result.returncode == 0:
            pages = result.stdout.decode("utf-8", "replace").split("\f")
            if pages and not pages[-1].strip():
                pages = pages[:-1]
            return Conversion(pages, "pdftotext")
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            return Conversion([p.extract_text() or "" for p in reader.pages], "pypdf")
    except ImportError:
        return Conversion([], "none", ["No PDF tool: install poppler (pdftotext) or `pip install pypdf`."])
    except Exception as error:  # damaged file
        return Conversion([], "none", [f"Could not read PDF: {error}"])


def _via_pdf(path: Path) -> Conversion:
    """Office → PDF with LibreOffice first: direct text export loses formulas."""
    out = tempfile.mkdtemp()  # the build moves the PDF into the hidden folder (Conversion.view)
    result = subprocess.run([_soffice(), "--headless", "--convert-to", "pdf", "--outdir", out, str(path)],
                            capture_output=True, timeout=300)
    pdf = Path(out) / (path.stem + ".pdf")
    if result.returncode == 0 and pdf.exists():
        conversion = _pdf_pages(pdf)
        conversion.method = "soffice+" + conversion.method
        conversion.view = pdf
        return conversion
    shutil.rmtree(out, ignore_errors=True)
    return None


def _docx_text(path: Path) -> Conversion:
    import docx
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return Conversion(["\n".join(parts)], "python-docx",
                      ["Formulas may be missing: install LibreOffice for full fidelity."])


def _markitdown(path: Path):
    """Microsoft MarkItDown (pip install 'markitdown[all]'), when installed: many formats, one page."""
    try:
        from markitdown import MarkItDown
    except ImportError:
        return None
    try:
        text = MarkItDown().convert(str(path)).text_content
    except Exception as error:
        return Conversion([], "none", [f"markitdown could not read it: {error}"])
    return Conversion([text], "markitdown")


def _pptx_text(path: Path) -> Conversion:
    from pptx import Presentation
    slides = []
    for slide in Presentation(str(path)).slides:
        texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        slides.append("\n".join(texts))
    return Conversion(slides, "python-pptx")


def _pictures(path) -> dict:
    """{slide number: its pictures} of a .pptx; empty without python-pptx."""
    if Path(path).suffix.lower() != ".pptx":
        return {}
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        return {}
    return {number: pictures for number, slide in enumerate(Presentation(str(path)).slides, 1)  # ponytail: top-level
            if (pictures := [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE])}  # pictures only


def slide_pictures(path, pages, folder) -> list:
    """The pictures on these slides of a .pptx, saved as image files a worker can open (Read can't open a .pptx)."""
    found, out = _pictures(path), []
    for number in pages:
        for k, shape in enumerate(found.get(number, []), 1):
            target = Path(folder) / f"slide-{number}-{k}.{shape.image.ext}"
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(shape.image.blob)
            out.append(str(target))
    return out


ANSWER_LINES = re.compile(r"([^\w\s]|_)\1{2,}")  # a question's ____ lines, a contents page's .... leaders
WORD = re.compile(r"[^\W\d_]+|\d+")
EAST_ASIAN = {"CJK", "HIRAGANA", "KATAKANA", "HANGUL", "IDEOGRAPHIC"}  # written without spaces, several scripts in one run


def _script(char):
    name = unicodedata.name(char, "?").split()[0].split("-")[0]  # KATAKANA-HIRAGANA PROLONGED SOUND MARK
    return "CJK" if name in EAST_ASIAN else name


def is_noise(text) -> bool:
    """A text layer that is OCR noise, as a handwritten scan's ("'fr Eg- I d 4 Lh tL"): under half of its characters
    are in numbers or words (2+ letters of one script; 3+ for Latin, since noise is short Latin runs). Bidi marks and
    vowel points are left out first: a Hebrew table is full of them."""
    text = "".join(c for c in text if unicodedata.category(c) not in ("Cf", "Mn"))
    text = ANSWER_LINES.sub(" ", text)
    total = sum(not c.isspace() for c in text)
    words = 0
    for word in WORD.findall(text):
        scripts = {_script(c) for c in word}
        if word[0].isdigit() or (len(scripts) == 1 and len(word) >= (3 if scripts == {"LATIN"} else 2)):
            words += len(word)
    return words < total / 2


def convert(path) -> Conversion:
    conversion = _convert(Path(path))
    if conversion.method not in NOT_PAGED:  # an unreadable text layer is no text layer: the page is read visually
        conversion.pages = ["" if is_noise(text) else text for text in conversion.pages]
    if conversion.pages:
        conversion.picture_pages = sorted(n for n in _pictures(path) if n <= len(conversion.pages))
    return conversion


def _slides(conversion):
    """MarkItDown's .pptx text, one page per slide (it marks each slide), so pages can be cited and read."""
    parts = re.split(r"<!-- Slide number: \d+ -->\n?", conversion.pages[0])
    return Conversion([p.strip() for p in parts[1:]], "markitdown-slides") if len(parts) > 1 else conversion


def _convert(path) -> Conversion:
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        return Conversion([path.read_text("utf-8", errors="replace")], "text")
    if suffix == ".pdf":
        return _pdf_pages(path)  # page by page, so citations can point at a page
    if suffix in MARKITDOWN_ONLY:
        return _markitdown(path) or Conversion([], "none", [f"Install markitdown to read {suffix} files."])
    if suffix in OFFICE:
        if _soffice():  # LibreOffice → PDF keeps formulas and pages
            conversion = _via_pdf(path)
            if conversion is not None:
                return conversion
        conversion = _markitdown(path)
        if conversion is not None and conversion.pages:
            return _slides(conversion) if suffix == ".pptx" else conversion
        try:
            if suffix == ".docx":
                return _docx_text(path)
            if suffix == ".pptx":
                return _pptx_text(path)
        except ImportError:
            pass
        return Conversion([], "none", [f"Cannot convert {suffix}: install LibreOffice."])
    return Conversion([], "none", [f"Not a document: {suffix}"])
